from pathlib import Path

from openpyxl import load_workbook

from photoacoustic.analysis import compute_alpha
from photoacoustic.constants import OPTIONS
from photoacoustic.models import Experiment, MeasurementFile, PowerScan
import pandas as pd


def reorganize_sheets(path: Path):
    """Reorganize sheets in an excel file"""
    wb = load_workbook(path)
    sheetnames = wb.sheetnames
    for ndx, sheetname in enumerate(sheetnames, 0):
        wb.move_sheet(sheetname, -ndx)

    sheetnames = wb.sheetnames
    for ndx, sheetname in enumerate(sheetnames, 0):
        if sheetname.startswith("ref") or sheetname.startswith("sam"):
            wb.move_sheet(sheetname, -ndx + 2)

    wb.save(path)


def write_excel(root: Path, exp: Experiment):
    excel_path = root / "summary.xlsx"
    with pd.ExcelWriter(excel_path) as xlsx:
        for _, powerscan in exp.powerscans.items():
            for _, measurement_file in powerscan.measurement_files.items():
                write_file_sheet(xlsx, measurement_file)
            write_powerscan_sheet(xlsx, powerscan)
        result = compute_alpha(exp)
        if result is not None:
            result.to_excel(
                xlsx, sheet_name="__ALPHA__", startrow=0, index=False, header=True
            )
        else:
            OPTIONS["on_error"]("failed to create results excel")

    reorganize_sheets(excel_path)


def write_file_sheet(xlsx: pd.ExcelWriter, measurement_file: MeasurementFile) -> None:
    path = Path(measurement_file.metadata.PATH)
    try:
        prefix = "(%s)" % path.parent.stem.split("_")[0]
    except Exception:
        prefix = "(?)"
    measurement_file.to_pandas().to_excel(
        xlsx, sheet_name=prefix + " " + path.stem, startrow=0, index=False, header=True
    )


def write_powerscan_sheet(xlsx: pd.ExcelWriter, powerscan: PowerScan) -> None:
    powerscan.to_pandas().to_excel(xlsx, sheet_name=powerscan.path.stem, index=False)
